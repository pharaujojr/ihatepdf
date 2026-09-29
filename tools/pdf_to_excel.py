#!/usr/bin/env python3
"""Converte as tabelas de um PDF em uma planilha .xlsx.

- Detecção: PyMuPDF find_tables() — primeiro pelas linhas desenhadas (tabelas
  com grade); se a página não tiver nenhuma, tenta pelo alinhamento do texto
  (tabelas sem borda, tipo extrato bancário), com filtros contra falsos
  positivos.
- Tabelas que continuam na página seguinte (mesmo nº de colunas, no topo da
  página ou com o cabeçalho repetido) são emendadas numa só.
- Números no padrão brasileiro (1.234,56 / R$ / %) viram números de verdade
  e datas dd/mm/aaaa viram datas, para dar pra somar e filtrar no Excel.
- Se nenhuma tabela for encontrada, gera uma aba "Texto" com as linhas do PDF
  divididas em colunas pelos espaçamentos.

Uso: pdf_to_excel.py <entrada.pdf> <saida.xlsx> [--single-sheet] [--no-merge] [--raw]
Imprime um JSON com estatísticas no stdout.
"""
import datetime
import json
import re
import sys

import pymupdf as fitz
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill('solid', fgColor='8B0000')
HEADER_FONT = Font(bold=True, color='FFFFFF')
TITLE_FONT = Font(bold=True, size=12, color='8B0000')
MAX_COL_WIDTH = 60

NUM_BR = re.compile(r'^\(?([-+]?)\s*(R\$|US\$|\$|€)?\s*([-+]?)(\d{1,3}(?:\.\d{3})+|\d+)(,\d+)?\)?\s*(%|[DC])?$')
NUM_US = re.compile(r'^\(?([-+]?)\s*(US\$|\$)?\s*([-+]?)(\d{1,3}(?:,\d{3})+)(\.\d+)?\)?$')
DATE_RE = re.compile(r'^(\d{1,2})/(\d{1,2})/(\d{2}|\d{4})$')
DATETIME_RE = re.compile(r'^(\d{1,2})/(\d{1,2})/(\d{4})\s+(\d{1,2}):(\d{2})(?::(\d{2}))?$')


# --------------------------------------------------------------------------
# Conversão de células
# --------------------------------------------------------------------------
def convert_cell(raw):
    """Retorna (valor, number_format) — texto fica como está."""
    text = raw.strip()
    if not text or len(text) > 40:
        return text, None

    m = DATETIME_RE.match(text)
    if m:
        d, mo, y, hh, mm, ss = (int(g) if g else 0 for g in m.groups())
        try:
            return datetime.datetime(y, mo, d, hh, mm, ss), 'DD/MM/YYYY HH:MM:SS' if m.group(6) else 'DD/MM/YYYY HH:MM'
        except ValueError:
            return text, None

    m = DATE_RE.match(text)
    if m:
        d, mo, y = (int(g) for g in m.groups())
        if y < 100:
            y += 2000 if y < 70 else 1900
        try:
            return datetime.datetime(y, mo, d), 'DD/MM/YYYY'
        except ValueError:
            return text, None

    compact = text.replace(' ', ' ')
    m = NUM_BR.match(compact)
    us = False
    if not m:
        m = NUM_US.match(compact)
        us = bool(m)
    if not m:
        return text, None

    if us:
        sign1, cur, sign2, intpart, dec = m.groups()
        suffix = None
        number = intpart.replace(',', '') + (dec or '')
    else:
        sign1, cur, sign2, intpart, dec, suffix = m.groups()
        # "0012" é código (CEP, matrícula...), não número
        if len(intpart) > 1 and intpart.startswith('0') and not dec:
            return text, None
        number = intpart.replace('.', '') + ('.' + dec[1:] if dec else '')
    digits = re.sub(r'\D', '', number)
    if len(digits) > 15:  # além da precisão do Excel: mantém como texto
        return text, None
    try:
        value = float(number)
    except ValueError:
        return text, None
    negative = '-' in (sign1 + sign2) or (text.startswith('(') and text.endswith(')')) \
        or suffix == 'D'
    if negative:
        value = -value
    if value.is_integer() and not dec and not cur and suffix != '%':
        value = int(value)

    if suffix == '%':
        return value / 100, '0.00%'
    if cur:
        symbol = 'R$' if cur == 'R$' else cur
        return value, '"{0}" #,##0.00;-"{0}" #,##0.00'.format(symbol)
    if dec:
        places = min(len(dec) - 1, 6)
        return value, '#,##0.' + '0' * places if places else '#,##0'
    if isinstance(value, int) and abs(value) >= 10000 and '.' in intpart:
        return value, '#,##0'
    return value, None


def clean(cell):
    if cell is None:
        return ''
    text = str(cell).replace('\r', '')
    lines = [re.sub(r'[ \t ]+', ' ', l).strip() for l in text.split('\n')]
    return '\n'.join(l for l in lines if l)


def tidy(rows):
    """Remove linhas e colunas totalmente vazias e junta colunas fantasmas."""
    rows = [[clean(c) for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return rows
    ncols = max(len(r) for r in rows)
    rows = [r + [''] * (ncols - len(r)) for r in rows]
    keep = [j for j in range(ncols) if any(r[j] for r in rows)]
    rows = [[r[j] for j in keep] for r in rows]
    return collapse_columns(rows)


def collapse_columns(rows):
    """Células mescladas no PDF espalham os valores por várias colunas da grade.
    Duas colunas vizinhas que nunca têm valor na mesma linha são a mesma coluna
    — exceto quando ambas têm cabeçalho próprio (ex.: Débito | Crédito)."""
    if not rows:
        return rows
    cols = [list(c) for c in zip(*rows)]
    j = 0
    while j < len(cols) - 1:
        a, b = cols[j], cols[j + 1]
        both_headed = a[0] and b[0]
        clash = any(x and y for x, y in zip(a, b))
        if not clash and not both_headed:
            cols[j] = [x or y for x, y in zip(a, b)]
            del cols[j + 1]
        else:
            j += 1
    return [list(r) for r in zip(*cols)]


def fill_ratio(rows):
    cells = [c for r in rows for c in r]
    return sum(1 for c in cells if c) / max(len(cells), 1)


# --------------------------------------------------------------------------
# Detecção de tabelas
# --------------------------------------------------------------------------
def extract_table(tab):
    rows = tab.extract()
    try:
        header = tab.header
        if header.external and header.names:
            rows = [list(header.names)] + rows
    except Exception:
        pass
    return tidy(rows)


def plausible_grid_table(rows):
    """Tabelas com grade: descarta desenhos técnicos/molduras que parecem grade."""
    if len(rows) < 2 or len(rows[0]) < 2:
        return False
    return fill_ratio(rows) >= 0.4


def plausible_text_table(rows):
    """Filtro para a estratégia 'text', que acha tabela até em parágrafo e sumário."""
    if len(rows) < 3 or len(rows[0]) < 2:
        return False
    if fill_ratio(rows) < 0.6:
        return False
    cells = [c for r in rows for c in r if c]
    # Tabela de verdade tem células curtas; prosa tem células enormes
    if sum(len(c) for c in cells) / len(cells) > 40:
        return False
    # Sumário: pontilhados guia ("Capítulo ........ 3")
    if sum(1 for c in cells if '....' in c) > len(cells) * 0.1:
        return False
    # Palavra cortada entre colunas ("Cont" | "role") = colunas falsas
    pairs = split = 0
    for r in rows:
        for a, b in zip(r, r[1:]):
            if a and b:
                pairs += 1
                if a[-1].isalpha() and b[0].islower():
                    split += 1
    return not pairs or split / pairs < 0.1


def find_page_tables(page):
    found = []
    try:
        for tab in page.find_tables().tables:
            rows = extract_table(tab)
            if plausible_grid_table(rows):
                found.append({'rows': rows, 'bbox': tuple(tab.bbox), 'strategy': 'lines'})
    except Exception:
        pass
    if found:
        return found
    return find_borderless_tables(page)


# --------------------------------------------------------------------------
# Tabelas sem grade: colunas pelos "rios" de espaço vazio entre as palavras
# --------------------------------------------------------------------------
def word_rows(words):
    """Agrupa palavras (x0, y0, x1, y1, texto, ...) em linhas visuais."""
    rows = []
    for w in sorted(words, key=lambda w: ((w[1] + w[3]) / 2, w[0])):
        yc, h = (w[1] + w[3]) / 2, w[3] - w[1]
        if rows and abs(rows[-1]['yc'] - yc) < max(rows[-1]['h'], h) * 0.5:
            rows[-1]['words'].append(w)
        else:
            rows.append({'yc': yc, 'h': h, 'words': [w]})
    for r in rows:
        r['words'].sort(key=lambda w: w[0])
        r['y0'] = min(w[1] for w in r['words'])
        r['y1'] = max(w[3] for w in r['words'])
    return rows


def segments(row):
    """Trechos da linha separados por vão maior que ~1 caractere de altura."""
    segs, cur = [], [row['words'][0]]
    for w in row['words'][1:]:
        if w[0] - cur[-1][2] > row['h'] * 0.8:
            segs.append(cur)
            cur = [w]
        else:
            cur.append(w)
    segs.append(cur)
    return segs


def column_bounds(rows, min_gap):
    """Faixas de x onde nenhuma palavra passa = separadores de coluna."""
    spans = sorted((w[0], w[2]) for r in rows for w in r['words'])
    merged = []
    for x0, x1 in spans:
        if merged and x0 - merged[-1][1] < min_gap:
            merged[-1][1] = max(merged[-1][1], x1)
        else:
            merged.append([x0, x1])
    return [(a[1] + b[0]) / 2 for a, b in zip(merged, merged[1:])]


def build_rows(rows, cuts):
    out = []
    for r in rows:
        cells = [[] for _ in range(len(cuts) + 1)]
        for w in r['words']:
            xc = (w[0] + w[2]) / 2
            cells[sum(1 for c in cuts if xc > c)].append(w[4])
        out.append([' '.join(c) for c in cells])
    return out


def find_borderless_tables(page):
    r = page.rect
    # Fora das margens: não engole cabeçalho/rodapé da página como linha da tabela
    clip = fitz.Rect(r.x0, r.y0 + r.height * 0.07, r.x1, r.y1 - r.height * 0.07)
    rows = word_rows(page.get_text('words', clip=clip))
    runs, cur = [], []
    for row in rows:
        multi = len(segments(row)) >= 2
        close = cur and row['y0'] - cur[-1]['y1'] < max(row['h'], cur[-1]['h']) * 1.8
        if multi and (not cur or close):
            cur.append(row)
        else:
            if len(cur) >= 3:
                runs.append(cur)
            cur = [row] if multi else []
    if len(cur) >= 3:
        runs.append(cur)

    found = []
    for run in runs:
        seg_counts = sorted(len(segments(x)) for x in run)
        if seg_counts[len(seg_counts) // 2] < 3:  # mediana: pelo menos 3 colunas
            continue
        heights = sorted(x['h'] for x in run)
        cuts = column_bounds(run, heights[len(heights) // 2] * 0.8)
        table = tidy(build_rows(run, cuts))
        if plausible_text_table(table):
            bbox = (min(w[0] for x in run for w in x['words']), run[0]['y0'],
                    max(w[2] for x in run for w in x['words']), run[-1]['y1'])
            found.append({'rows': table, 'bbox': bbox, 'strategy': 'text'})
    return found


def same_row(a, b):
    return [c.lower() for c in a] == [c.lower() for c in b]


def collect_tables(doc, merge):
    tables = []
    for pno, page in enumerate(doc):
        page_h = page.rect.height
        for idx, t in enumerate(find_page_tables(page)):
            rows = t['rows']
            prev = tables[-1] if tables else None
            if merge and prev and idx == 0 and prev['last_page'] == pno - 1 \
                    and len(prev['rows'][0]) == len(rows[0]):
                repeats_header = same_row(rows[0], prev['rows'][0])
                at_top = t['bbox'][1] < page_h * 0.25
                if repeats_header or at_top:
                    prev['rows'].extend(rows[1:] if repeats_header else rows)
                    prev['last_page'] = pno
                    continue
            tables.append({'rows': rows, 'first_page': pno, 'last_page': pno})
    return tables


def text_fallback(doc):
    """Sem tabelas: cada linha do PDF vira uma linha da planilha, com colunas
    separadas onde há vão horizontal grande entre as palavras."""
    rows = []
    for pno, page in enumerate(doc):
        words = page.get_text('words')  # x0, y0, x1, y1, texto, bloco, linha, nº
        lines = {}
        for w in words:
            lines.setdefault((w[5], w[6]), []).append(w)
        ordered = sorted(lines.values(), key=lambda ws: (round(ws[0][1]), ws[0][0]))
        for ws in ordered:
            ws.sort(key=lambda w: w[0])
            cells, current, last_x1 = [], [], None
            for w in ws:
                height = w[3] - w[1]
                if last_x1 is not None and w[0] - last_x1 > height * 1.2:
                    cells.append(' '.join(current))
                    current = []
                current.append(w[4])
                last_x1 = w[2]
            if current:
                cells.append(' '.join(current))
            rows.append(cells)
        if pno < doc.page_count - 1:
            rows.append([])
    return rows


# --------------------------------------------------------------------------
# Escrita do .xlsx
# --------------------------------------------------------------------------
def write_rows(ws, rows, start_row, raw, header=True):
    widths = {}
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            if raw:
                value, fmt = cell, None
            else:
                value, fmt = convert_cell(cell) if not (header and i == 0) else (cell, None)
            c = ws.cell(row=start_row + i, column=j + 1, value=value)
            if fmt:
                c.number_format = fmt
            if isinstance(value, str) and '\n' in value:
                c.alignment = Alignment(wrap_text=True, vertical='top')
            if header and i == 0:
                c.fill = HEADER_FILL
                c.font = HEADER_FONT
                c.alignment = Alignment(wrap_text=True, vertical='center')
            longest = max((len(p) for p in str(cell).split('\n')), default=0)
            widths[j] = max(widths.get(j, 0), longest)
    for j, w in widths.items():
        letter = get_column_letter(j + 1)
        current = ws.column_dimensions[letter].width or 0
        ws.column_dimensions[letter].width = max(current, min(MAX_COL_WIDTH, w + 2))
    return start_row + len(rows)


def page_label(t):
    if t['first_page'] == t['last_page']:
        return 'p.{0}'.format(t['first_page'] + 1)
    return 'p.{0}-{1}'.format(t['first_page'] + 1, t['last_page'] + 1)


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    opts = {a for a in sys.argv[1:] if a.startswith('--')}
    if len(args) < 2:
        print(__doc__, file=sys.stderr)
        sys.exit(2)
    src, dst = args[0], args[1]
    # O PyMuPDF imprime avisos no stdout; o JSON final precisa sair sozinho lá
    real_stdout, sys.stdout = sys.stdout, sys.stderr
    single = '--single-sheet' in opts
    merge = '--no-merge' not in opts
    raw = '--raw' in opts

    doc = fitz.open(src)
    if doc.needs_pass and not doc.authenticate(''):
        print('PDF protegido por senha.', file=sys.stderr)
        sys.exit(3)

    tables = collect_tables(doc, merge)
    wb = Workbook()
    stats = {'pages': doc.page_count, 'tables': len(tables), 'rows': 0, 'fallback': False}

    if not tables:
        rows = text_fallback(doc)
        if not any(rows):
            print('Nenhum texto encontrado (PDF digitalizado?).', file=sys.stderr)
            sys.exit(4)
        ws = wb.active
        ws.title = 'Texto'
        write_rows(ws, rows, 1, raw, header=False)
        stats['fallback'] = True
        stats['rows'] = len(rows)
    elif single:
        ws = wb.active
        ws.title = 'Tabelas'
        r = 1
        for n, t in enumerate(tables, 1):
            title = ws.cell(row=r, column=1, value='Tabela {0} ({1})'.format(n, page_label(t)))
            title.font = TITLE_FONT
            r = write_rows(ws, t['rows'], r + 1, raw) + 1
            stats['rows'] += len(t['rows'])
    else:
        wb.remove(wb.active)
        for n, t in enumerate(tables, 1):
            ws = wb.create_sheet(('Tabela {0} ({1})'.format(n, page_label(t)))[:31])
            write_rows(ws, t['rows'], 1, raw)
            ws.freeze_panes = 'A2'
            if len(t['rows']) > 1:
                ws.auto_filter.ref = 'A1:{0}{1}'.format(
                    get_column_letter(len(t['rows'][0])), len(t['rows']))
            stats['rows'] += len(t['rows'])

    wb.save(dst)
    print(json.dumps(stats, ensure_ascii=False), file=real_stdout)


if __name__ == '__main__':
    main()
